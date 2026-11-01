"""复盘报告导出：生成 Word (.docx)。

纯文本 + 表格版式（雷达图属可视化，Word 里用五维得分表替代），供打印/存档/批注。
只读取已生成的复盘，不涉及模型调用；密钥与评价逻辑仍全在服务端。
"""
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

_FONT = "Microsoft YaHei"
_ACCENT = RGBColor(0x07, 0x47, 0x2F)
_MUTED = RGBColor(0x38, 0x44, 0x3D)


def _style(run, size=None, bold=False, color=None):
    run.font.name = _FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), _FONT)
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _heading(doc, text, level=1, size=None):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        _style(
            run,
            size=size or (16 if level == 1 else 13),
            bold=True,
            color=_ACCENT if level == 1 else _MUTED,
        )
    return h


def _para(doc, text, size=10.5, bold=False, color=None):
    p = doc.add_paragraph()
    _style(p.add_run(text), size=size, bold=bold, color=color)
    return p


def _split_items(text):
    lines = [s.strip() for s in (text or "").replace("\r", "").split("\n") if s.strip()]
    if lines:
        return lines
    return [text.strip()] if text and text.strip() else []


def build_review_docx(evaluation, turns, session_id: str, prep_note: str = "") -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = _FONT
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), _FONT)
    normal.font.size = Pt(10.5)

    dims = evaluation.dimensions or []

    _heading(doc, "谈话训练评价报告", level=1, size=18)

    # 基本信息
    mode_label = ""
    for t in turns:
        if t.role == "teacher":
            mode_label = "实时通话" if t.mode == "realtime" else "语音/文字对练"
            break
    facts = [
        ("训练场景", evaluation.scenario or "—"),
        ("训练方式", mode_label or "—"),
        ("综合得分", f"{evaluation.total_score} / 100"),
        ("评价维度", f"{len(dims)} 维"),
        ("会话编号", (session_id or "")[:8] or "—"),
    ]
    for k, v in facts:
        p = doc.add_paragraph()
        _style(p.add_run(f"{k}："), size=10.5, bold=True)
        _style(p.add_run(v), size=10.5)

    # 谈前准备（教师可选择是否纳入，未纳入时 prep_note 传空即整段不生成）
    if prep_note and prep_note.strip():
        _heading(doc, "谈前准备")
        for line in _split_items(prep_note):
            _para(doc, line)

    _heading(doc, "一、整体评价")
    _para(doc, evaluation.summary or "（无）")

    _heading(doc, "二、五维评价")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, htext in enumerate(["维度", "得分", "点评"]):
        _style(table.rows[0].cells[i].paragraphs[0].add_run(htext), size=10.5, bold=True)
    for d in dims:
        cells = table.add_row().cells
        vals = [d.get("name", ""), str(d.get("score", 0)), d.get("comment", "") or ""]
        for i, val in enumerate(vals):
            _style(cells[i].paragraphs[0].add_run(val), size=10.5)

    _heading(doc, "三、完整对话记录")
    for t in turns:
        who = "辅导员" if t.role == "teacher" else "学生"
        p = doc.add_paragraph()
        _style(
            p.add_run(f"{who}："),
            size=10.5,
            bold=True,
            color=_ACCENT if t.role == "teacher" else _MUTED,
        )
        _style(p.add_run(t.text or ""), size=10.5)

    _heading(doc, "四、五维点评与证据")
    for d in dims:
        _heading(doc, f"{d.get('name', '')} · {d.get('score', 0)} 分", level=2)
        _para(doc, d.get("comment", "") or "（无点评）")
        evidence = [q for q in (d.get("evidence") or []) if q]
        if evidence:
            _para(doc, "原话依据：", size=10.5, bold=True)
            for q in evidence:
                _para(doc, f"“{q}”")

    _heading(doc, "五、下一轮练习计划")
    items = _split_items(evaluation.suggestions)
    if not items:
        _para(doc, "暂无明确建议。")
    for i, s in enumerate(items, 1):
        _para(doc, f"{i}. {s}")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
